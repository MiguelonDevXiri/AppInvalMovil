import json
import os
import re
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

EXCEL = Path('/home/miguel/.openclaw/media/inbound/filtros_actualizado---dd68618f-3de6-4e5f-a7ee-a6a72f511ed1.xlsx')
ROOT = Path.cwd()


def load_env() -> None:
    for envp in [
        Path('/home/miguel/Documents/Mis proyectos /Prroyecto APP/inval-panel/.env.local'),
        ROOT / '.env',
        ROOT / '.env.local',
    ]:
        if not envp.exists():
            continue
        for line in envp.read_text().splitlines():
            line = line.strip()
            if not line or line.startswith('#') or '=' not in line:
                continue
            key, value = line.split('=', 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"'))


def norm_cell(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def norm_ref(value):
    text = norm_cell(value).upper().strip()
    text = re.sub(r'\s+', ' ', text)
    return text.strip(' /')


def qty(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0
    try:
        return int(float(str(value).replace(',', '.')))
    except Exception:
        return 0


def opt_qty(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        parsed = int(float(str(value).replace(',', '.')))
        return parsed if parsed >= 0 else None
    except Exception:
        return None


def split_codes(value):
    raw = norm_cell(value).upper()
    if not raw:
        return []
    parts = [part.strip() for part in re.split(r'[/,]', raw) if part.strip()]
    if all(re.fullmatch(r'0?\d{5,8}', part) for part in parts):
        return parts
    return []


def clean_alternative(value):
    value = norm_ref(value)
    # Quita precios anotados al final: "WD10018 21,58" -> "WD10018".
    value = re.sub(r'\s+\d+[,.]\d+\s*$', '', value).strip()
    return value


def split_ref_values(value):
    raw = norm_cell(value).upper()
    if not raw:
        return []
    raw = re.sub(r'\s+-\s+', ' / ', raw)
    raw = re.sub(r'\s+/\s+', ' / ', raw)
    raw = re.sub(r'(?<=[A-Z0-9])/(?=[A-Z]+\d)', ' / ', raw)
    parts = []
    for part in re.split(r'\s+/\s+|;', raw):
        cleaned = clean_alternative(part)
        if cleaned and '???' not in cleaned:
            parts.append(cleaned)
    return parts


def is_note(reference):
    reference = norm_ref(reference)
    if not reference:
        return True
    return any(word in reference for word in ['VER FOTOS', 'SE PUEDE', 'FILTRO TDC', 'PEDIR PARA', '-->'])


def parse_excel():
    wb = load_workbook(EXCEL, data_only=True)
    ws = wb['stock filtros']
    rows = list(ws.iter_rows(values_only=True))
    header = [norm_cell(cell).lower() for cell in rows[1]]
    idx = {name: index for index, name in enumerate(header)}

    stock = {}
    crosses = []
    seen = set()
    skipped = []

    def add_cross(reference, equivalent, brand):
        equivalent = norm_ref(equivalent)
        if not equivalent or equivalent == reference:
            return
        key = (reference, equivalent)
        if key in seen:
            return
        seen.add(key)
        crosses.append({
            'reference': reference,
            'equivalent_reference': equivalent,
            'brand': brand,
            'notes': 'Excel filtros actualizado 2026-06-05',
        })

    for row_no, row in enumerate(rows[2:], start=3):
        reference = norm_ref(row[idx['referencia']])
        if is_note(reference):
            if any(norm_cell(cell) for cell in row):
                skipped.append({'row': row_no, 'reference': reference, 'reason': 'sin referencia/nota'})
            continue

        minimum = opt_qty(row[idx['deberia']])
        if minimum is None:
            minimum = 2

        stock[reference] = {
            'reference': reference,
            'description': norm_cell(row[idx['uso']]) or None,
            'warehouse_qty': qty(row[idx['almacen']]),
            'van_qty': qty(row[idx['furgoneta']]),
            'minimum_qty': minimum,
            'recommended_qty': max(4, minimum),
        }

        for code in split_codes(row[idx['código']]):
            add_cross(reference, code, 'Código INVAL')
        for alternative in split_ref_values(row[idx['alternativa']]):
            add_cross(reference, alternative, 'Alternativa 1')
        for alternative in split_ref_values(row[idx['alternativa 2']]):
            add_cross(reference, alternative, 'Alternativa 2')

    return list(stock.values()), crosses, skipped


load_env()
SUPABASE_URL = os.environ.get('NEXT_PUBLIC_SUPABASE_URL') or os.environ.get('EXPO_PUBLIC_SUPABASE_URL')
SUPABASE_KEY = os.environ.get('NEXT_PUBLIC_SUPABASE_ANON_KEY') or os.environ.get('EXPO_PUBLIC_SUPABASE_ANON_KEY')
if not SUPABASE_URL or not SUPABASE_KEY:
    raise SystemExit('Faltan credenciales de Supabase')

HEADERS = {
    'apikey': SUPABASE_KEY,
    'Authorization': f'Bearer {SUPABASE_KEY}',
    'Accept': 'application/json',
    'Content-Type': 'application/json',
}


def request(method, path, payload=None, extra_headers=None):
    headers = dict(HEADERS)
    if extra_headers:
        headers.update(extra_headers)
    data = json.dumps(payload).encode() if payload is not None else None
    req = urllib.request.Request(f'{SUPABASE_URL}/rest/v1/{path}', data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req) as response:
            body = response.read().decode()
            return json.loads(body) if body else None
    except urllib.error.HTTPError as error:
        body = error.read().decode()
        raise RuntimeError(f'{method} {path} -> {error.code}: {body}') from error


def get_all(table):
    return request('GET', f'{table}?select=*') or []


stock, crosses, skipped = parse_excel()
backup = {table: get_all(table) for table in ['filters_stock', 'filters_crosses', 'filters_movements', 'filters_stock_counts']}
backup_path = ROOT / 'tmp' / 'filters_sync_updated_backup_20260605_1226.json'
backup_path.write_text(json.dumps(backup, ensure_ascii=False, indent=2))

# Update catalog, preserve movement history and count history.
# The public REST key cannot delete old crosses because there is no DELETE RLS policy,
# so crosses are upserted. Obsolete old crosses require SQL/service access to prune.
for index in range(0, len(stock), 50):
    request(
        'POST',
        'filters_stock',
        stock[index:index + 50],
        {'Prefer': 'resolution=merge-duplicates,return=minimal'},
    )
for index in range(0, len(crosses), 50):
    request(
        'POST',
        'filters_crosses?on_conflict=reference,equivalent_reference',
        crosses[index:index + 50],
        {'Prefer': 'resolution=merge-duplicates,return=minimal'},
    )

summary = {
    'backup': str(backup_path),
    'stock_loaded': len(stock),
    'crosses_loaded': len(crosses),
    'skipped_count': len(skipped),
    'current_stock_count': len(get_all('filters_stock')),
    'current_cross_count': len(get_all('filters_crosses')),
    'samples': stock[:5],
}
summary_path = ROOT / 'tmp' / 'filters_sync_updated_summary_20260605_1226.json'
summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2))
print(json.dumps(summary, ensure_ascii=False, indent=2))

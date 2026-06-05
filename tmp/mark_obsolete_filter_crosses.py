import json
import os
import re
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from openpyxl import load_workbook

EXCEL = Path('/home/miguel/.openclaw/media/inbound/filtros_actualizado---dd68618f-3de6-4e5f-a7ee-a6a72f511ed1.xlsx')
ROOT = Path.cwd()
for envp in [Path('/home/miguel/Documents/Mis proyectos /Prroyecto APP/inval-panel/.env.local'), ROOT / '.env', ROOT / '.env.local']:
    if envp.exists():
        for line in envp.read_text().splitlines():
            line=line.strip()
            if line and not line.startswith('#') and '=' in line:
                k,v=line.split('=',1); os.environ.setdefault(k.strip(), v.strip().strip('"'))
URL=os.environ.get('NEXT_PUBLIC_SUPABASE_URL') or os.environ.get('EXPO_PUBLIC_SUPABASE_URL')
KEY=os.environ.get('NEXT_PUBLIC_SUPABASE_ANON_KEY') or os.environ.get('EXPO_PUBLIC_SUPABASE_ANON_KEY')
H={'apikey':KEY,'Authorization':'Bearer '+KEY,'Accept':'application/json','Content-Type':'application/json'}

def req(method,path,payload=None):
    data=json.dumps(payload).encode() if payload is not None else None
    r=urllib.request.Request(f'{URL}/rest/v1/{path}',data=data,headers=H,method=method)
    try:
        with urllib.request.urlopen(r) as resp:
            body=resp.read().decode(); return json.loads(body) if body else None
    except urllib.error.HTTPError as e:
        raise RuntimeError(f'{method} {path} {e.code}: {e.read().decode()}')

def norm_cell(v):
    if v is None: return ''
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()
def norm_ref(v):
    s=norm_cell(v).upper().strip(); s=re.sub(r'\s+',' ',s); return s.strip(' /')
def split_codes(v):
    raw=norm_cell(v).upper(); parts=[p.strip() for p in re.split(r'[/,]',raw) if p.strip()]
    return parts if parts and all(re.fullmatch(r'0?\d{5,8}',p) for p in parts) else []
def clean_alternative(v):
    v=norm_ref(v)
    return re.sub(r'\s+\d+[,.]\d+\s*$', '', v).strip()

def split_ref_values(v):
    raw=norm_cell(v).upper()
    if not raw: return []
    raw=re.sub(r'\s+-\s+', ' / ', raw)
    raw=re.sub(r'\s+/\s+', ' / ', raw)
    raw=re.sub(r'(?<=[A-Z0-9])/(?=[A-Z]+\d)', ' / ', raw)
    return [clean_alternative(p) for p in re.split(r'\s+/\s+|;', raw) if clean_alternative(p) and '???' not in clean_alternative(p)]
def is_note(ref):
    ref=norm_ref(ref)
    return (not ref) or any(w in ref for w in ['VER FOTOS','SE PUEDE','FILTRO TDC','PEDIR PARA','-->'])
wb=load_workbook(EXCEL,data_only=True); ws=wb['stock filtros']; rows=list(ws.iter_rows(values_only=True)); header=[norm_cell(x).lower() for x in rows[1]]; idx={n:i for i,n in enumerate(header)}
new=set()
def add(ref,eq):
    eq=norm_ref(eq)
    if eq and eq!=ref: new.add((ref,eq))
for row in rows[2:]:
    ref=norm_ref(row[idx['referencia']])
    if is_note(ref): continue
    for c in split_codes(row[idx['código']]): add(ref,c)
    for a in split_ref_values(row[idx['alternativa']]): add(ref,a)
    for a in split_ref_values(row[idx['alternativa 2']]): add(ref,a)
current=req('GET','filters_crosses?select=*') or []
obsolete=[c for c in current if (c['reference'], c['equivalent_reference']) not in new]
for c in obsolete:
    req('PATCH', 'filters_crosses?id=eq.' + urllib.parse.quote(c['id']), {
        'brand':'Obsoleto',
        'notes':'Obsoleto por Excel actualizado 2026-06-05',
    })
summary={'new_crosses':len(new),'current_crosses':len(current),'obsolete_marked':len(obsolete),'obsolete_samples':obsolete[:10]}
Path('tmp/obsolete_filter_crosses_summary_20260605_1226.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
print(json.dumps(summary,ensure_ascii=False,indent=2))
